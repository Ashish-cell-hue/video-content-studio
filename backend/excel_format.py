
import openpyxl
from openpyxl.utils import get_column_letter


def excel_color(color):
    """Convert Excel color to a browser-friendly value."""
    if color is None:
        return None

    if color.type == "rgb" and isinstance(color.rgb, str):
        rgb = color.rgb
        if len(rgb) == 8:
            return "#" + rgb[-6:]
        if len(rgb) == 6:
            return "#" + rgb

    return None


def get_sheet_format(excel_path, grade):
    """Read the original worksheet values and formatting."""
    allowed_grades = [f"Grade {n}" for n in range(6, 13)]

    if grade not in allowed_grades:
        raise ValueError("Invalid grade")

    workbook = openpyxl.load_workbook(
        excel_path,
        read_only=False,
        data_only=True
    )

    try:
        if grade not in workbook.sheetnames:
            raise ValueError("Grade sheet not found")

        sheet = workbook[grade]

        merged_ranges = [
            str(item) for item in sheet.merged_cells.ranges
        ]

        columns = {}
        for col in range(1, sheet.max_column + 1):
            letter = get_column_letter(col)
            dimension = sheet.column_dimensions[letter]
            columns[letter] = {
                "width": dimension.width,
                "hidden": dimension.hidden
            }

        rows = {}
        for row in range(1, sheet.max_row + 1):
            dimension = sheet.row_dimensions[row]
            rows[str(row)] = {
                "height": dimension.height,
                "hidden": dimension.hidden
            }

        cells = []

        for row in sheet.iter_rows():
            for cell in row:
                if cell.__class__.__name__ == "MergedCell":
                    continue

                value = cell.value

                if hasattr(value, "isoformat"):
                    value = value.isoformat()

                fill_color = None
                if cell.fill.patternType == "solid":
                    fill_color = excel_color(cell.fill.fgColor)

                cells.append({
                    "row": cell.row,
                    "col": cell.column,
                    "value": value,
                    "style": {
                        "bold": bool(cell.font.bold),
                        "italic": bool(cell.font.italic),
                        "fontSize": cell.font.sz,
                        "fontColor": excel_color(cell.font.color),
                        "background": fill_color,
                        "horizontal": cell.alignment.horizontal,
                        "vertical": cell.alignment.vertical,
                        "wrapText": bool(cell.alignment.wrap_text),
                        "border": {
                            side: getattr(cell.border, side).style
                            for side in (
                                "left", "right", "top", "bottom"
                            )
                        }
                    }
                })

        return {
            "sheet": grade,
            "maxRow": sheet.max_row,
            "maxCol": sheet.max_column,
            "merged": merged_ranges,
            "columns": columns,
            "rows": rows,
            "cells": cells
        }

    finally:
        workbook.close()
